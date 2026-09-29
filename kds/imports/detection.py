"""Deterministic, explainable assistance for supported structured imports.

Detection proposes a data type and mapping.  It never activates data and it
never fabricates a missing field.  Activation remains the responsibility of
the existing explicit-confirmation import contract.
"""
from __future__ import annotations

from collections import Counter
from io import BytesIO
from pathlib import PurePath
from typing import Any
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .mapping import FIELDS, REQUIRED, key, suggest
from .parser import ParseError, parse


LABELS = {
    "analysis_units": "Analiz birimleri", "crops": "Ürün kataloğu",
    "economics": "Genel ekonomi", "water_budget": "Su bütçesi",
    "geometries": "Geometri", "candidates": "Aday matrisi",
    "scientific_inputs": "Mevcut desen / iklim",
    "annual_water_supply": "Yıllık su arzı", "monthly_water_supply": "Aylık su arzı",
    "delivery_capacity": "Teslim kapasitesi", "environmental_release": "Çevresel akış",
    "conveyance_efficiency": "İletim randımanı",
    "perennial_irrigation_requirement": "Çok yıllık bitki su ihtiyacı",
    "crop_yield": "Ürün verimi", "crop_sale_price": "Satış fiyatı",
    "crop_support_payment": "Destek ödemesi", "crop_cost_components": "Maliyet bileşenleri",
    "crop_net_profit": "Net kâr", "analysis_unit_economics": "Birim ekonomisi",
    "seasonal_economics": "Mevsimsel ekonomi",
    "crop_water_parameters": "Ürün su parametreleri", "crop_phenology": "Fenoloji",
}

FILENAME_HINTS = {
    "analysis_units": ("analysis_unit", "analysis_units", "analiz_birim"),
    "crops": ("crops", "crop_catalog", "urun_katalog"),
    "economics": ("economics", "general_economics"),
    "water_budget": ("water_budget", "su_butce"),
    "geometries": ("geometry", "geometries", "geojson"),
    "candidates": ("candidate", "candidates", "aday"),
    "scientific_inputs": ("scientific", "current_pattern", "climate"),
    "annual_water_supply": ("annual_water", "annual_supply", "yillik_su"),
    "monthly_water_supply": ("monthly_water", "monthly_supply", "aylik_su"),
    "delivery_capacity": ("delivery_capacity", "teslim_kapas"),
    "environmental_release": ("environmental_release", "cevre_akis"),
    "conveyance_efficiency": ("conveyance_efficiency", "iletim_rand"),
    "perennial_irrigation_requirement": ("perennial_irrigation", "perennial_requirement"),
    "crop_yield": ("crop_yield", "urun_verim"),
    "crop_sale_price": ("crop_sale_price", "sale_price", "satis_fiyat"),
    "crop_support_payment": ("crop_support", "support_payment", "destek_odeme"),
    "crop_cost_components": ("crop_cost", "cost_components", "maliyet"),
    "crop_net_profit": ("crop_net_profit", "net_profit", "net_kar"),
    "analysis_unit_economics": ("analysis_unit_economics", "unit_economics"),
    "seasonal_economics": ("seasonal_economics", "mevsimsel_ekonomi"),
    "crop_water_parameters": ("crop_water_parameters", "water_parameters", "kc_parameters"),
    "crop_phenology": ("crop_phenology", "phenology", "fenoloji"),
}

# Generic structural expectations only. These fields are unambiguously
# numeric in the canonical import contracts; no agronomic threshold or
# scientific-value judgment is applied here.
NUMERIC_REQUIRED_FIELDS = {
    "area_da", "year", "planning_year", "applicable_year", "observation_year",
    "yield_per_da", "water_requirement_m3_da", "profit_per_da", "yield_ton_da",
    "amount", "capacity", "efficiency", "kc_ini", "kc_mid", "kc_end",
    "p_ini", "p_dev", "p_mid", "p_late", "yield_value", "price",
    "net_profit_per_da", "gross_revenue_per_da", "total_cost_per_da",
    "support_payment_per_da", "gross_revenue", "total_cost", "net_profit",
    "cost",
}
NUMERIC_REQUIRED_BY_TYPE = {
    "environmental_release": {"value"},
    "perennial_irrigation_requirement": {"value"},
}
YEAR_FIELDS = {"year", "planning_year", "applicable_year", "observation_year"}


def _workbook_sheets(content: bytes, extension: str) -> list[str]:
    if extension != ".xlsx":
        return []
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, TypeError, ValueError) as exc:
        raise ParseError(f"Invalid XLSX: {exc}") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _value_types(rows: list[dict[str, Any]], columns: list[str]) -> dict[str, dict[str, int]]:
    result = {}
    for column in columns:
        counts = Counter()
        for row in rows[:200]:
            value = row.get("values", {}).get(column)
            if value in (None, ""):
                counts["empty"] += 1
            elif isinstance(value, bool):
                counts["boolean"] += 1
            elif isinstance(value, (int, float)):
                counts["number"] += 1
            else:
                counts["text"] += 1
        result[column] = dict(counts)
    return result


def _mapped_values(rows: list[dict[str, Any]], mapping: dict[str, str], names: tuple[str, ...]) -> list[str]:
    values = []
    for name in names:
        source = mapping.get(name)
        if not source:
            continue
        for row in rows:
            value = row.get("values", {}).get(source)
            if value not in (None, "") and str(value).strip() not in values:
                values.append(str(value).strip())
    return values


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    supplied = str(value).strip().replace(" ", "")
    if "," in supplied and "." in supplied:
        if re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", supplied):
            supplied = supplied.replace(",", "")
        elif re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+(,\d+)?", supplied):
            supplied = supplied.replace(".", "").replace(",", ".")
        else:
            return None
    elif "," in supplied:
        supplied = supplied.replace(",", ".")
    try:
        return float(supplied)
    except ValueError:
        return None


def _structural_type_issues(rows: list[dict[str, Any]], mapping: dict[str, str],
                            required: list[str], data_type: str) -> list[dict[str, Any]]:
    issues = []
    numeric_fields = NUMERIC_REQUIRED_FIELDS | NUMERIC_REQUIRED_BY_TYPE.get(data_type, set())
    for canonical in required:
        source = mapping.get(canonical)
        if not source or canonical not in numeric_fields:
            continue
        populated = [row.get("values", {}).get(source) for row in rows
                     if row.get("values", {}).get(source) not in (None, "")]
        if not populated:
            continue
        parsed = [_number(value) for value in populated]
        invalid = sum(value is None for value in parsed)
        if invalid:
            invalid_examples = list(dict.fromkeys(
                str(value)[:120] for value, number in zip(populated, parsed) if number is None
            ))[:5]
            year = canonical in YEAR_FIELDS
            issues.append({
                "code": "explicit_year_not_parseable" if year else "required_numeric_field_not_parseable",
                "message": ("Explicit year field contains an unparseable numeric year."
                            if year else "A required numeric field contains a non-numeric value."),
                "field": canonical,
                "source_column": source,
                "populated": len(populated),
                "parseable": len(populated) - invalid,
                "unparseable": invalid,
                "invalid_examples": invalid_examples,
            })
    return issues


def detect(content: bytes, filename: str, options: dict[str, Any] | None = None,
           project: dict[str, Any] | None = None) -> dict[str, Any]:
    options = dict(options or {})
    extension = PurePath(filename).suffix.lower()
    base = {
        "filename": filename, "extension": extension, "state": "INVALID",
        "detected_data_type": None, "detected_domain": None, "mapping": {},
        "missing_required_fields": [], "ambiguous_mapping": {}, "issues": [],
        "rows": 0, "columns": [], "sheet_names": [], "selected_sheet": options.get("sheet"),
        "year": None, "scope": None, "authority": None, "confidence": 0.0,
    }
    if extension not in {".csv", ".xlsx", ".geojson"}:
        return {**base, "state": "UNSUPPORTED",
                "issues": [{"code": "unsupported_extension", "message": "CSV, XLSX or GeoJSON required."}]}
    try:
        sheets = _workbook_sheets(content, extension)
        sheet_evidence = []
        sheet_ambiguity = None
        if sheets and not options.get("sheet"):
            filename_key = key(PurePath(filename).stem)
            for sheet in sheets:
                try:
                    sheet_columns, sheet_rows = parse(content, filename, {**options, "sheet": sheet})
                except (ParseError, ValueError, OSError):
                    continue
                best_score = -1.0
                best_type = None
                for data_type in FIELDS:
                    mapping, ambiguous = suggest(sheet_columns, data_type)
                    required = REQUIRED[data_type]
                    coverage = sum(field in mapping for field in required) / max(1, len(required))
                    hint = any(key(token) in filename_key or key(token) in key(sheet)
                               for token in FILENAME_HINTS.get(data_type, ()))
                    score = coverage * 100 + (35 if hint else 0) - len(ambiguous) * 10
                    if score > best_score:
                        best_score, best_type = score, data_type
                sheet_evidence.append({"sheet": sheet, "rows": len(sheet_rows),
                                       "columns": len(sheet_columns), "best_type": best_type,
                                       "score": round(best_score, 3)})
            usable = [item for item in sheet_evidence if item["rows"] and item["columns"]]
            if not usable:
                raise ParseError("Workbook has no non-empty tabular sheet.")
            usable.sort(key=lambda item: (-item["score"], item["sheet"]))
            options["sheet"] = usable[0]["sheet"]
            tied_sheets = [item["sheet"] for item in usable if item["score"] == usable[0]["score"]]
            if len(tied_sheets) > 1:
                sheet_ambiguity = {"code": "ambiguous_workbook_sheet",
                                   "message": "Multiple workbook sheets have equal structural evidence.",
                                   "sheets": tied_sheets}
        columns, rows = parse(content, filename, options)
    except (ParseError, ValueError, OSError) as exc:
        return {**base, "state": "INVALID",
                "issues": [{"code": "invalid_file", "message": str(exc)}]}

    filename_key = key(PurePath(filename).stem)
    sheet_key = key(options.get("sheet", ""))
    candidates = []
    allowed = ("analysis_units", "geometries") if extension == ".geojson" else tuple(FIELDS)
    for data_type in allowed:
        mapping, ambiguous = suggest(columns, data_type)
        required = REQUIRED[data_type]
        matched = sum(field in mapping for field in required)
        coverage = matched / max(1, len(required))
        hints = FILENAME_HINTS.get(data_type, ())
        hint = any(key(token) in filename_key or key(token) in sheet_key for token in hints)
        exact_headers = sum(key(field) in {key(column) for column in columns} for field in required)
        score = coverage * 100 + exact_headers * 2 + (35 if hint else 0) - len(ambiguous) * 10
        candidates.append({"data_type": data_type, "mapping": mapping, "ambiguous": ambiguous,
                           "required": required, "matched": matched, "coverage": coverage,
                           "hint": hint, "score": score})
    candidates.sort(key=lambda item: (-item["score"], item["data_type"]))
    best = candidates[0]
    tied = [item for item in candidates if item["score"] == best["score"]]
    missing = [field for field in best["required"] if field not in best["mapping"]]
    issues = []
    if sheet_ambiguity:
        issues.append(sheet_ambiguity)
    if len(tied) > 1:
        issues.append({"code": "ambiguous_data_type", "message": "Multiple dataset types have equal evidence.",
                       "candidates": [item["data_type"] for item in tied]})
    if missing:
        issues.append({"code": "missing_required_fields", "message": "Required columns are missing.",
                       "fields": missing})
    if best["ambiguous"]:
        issues.append({"code": "ambiguous_mapping", "message": "One or more headers map ambiguously.",
                       "fields": best["ambiguous"]})
    type_issues = _structural_type_issues(rows, best["mapping"], best["required"], best["data_type"])
    issues.extend(type_issues)

    years = _mapped_values(rows, best["mapping"], ("planning_year", "applicable_year", "year", "observation_year"))
    scopes = _mapped_values(rows, best["mapping"], ("geographic_scope", "scope"))
    authorities = _mapped_values(rows, best["mapping"], ("authority_class",))
    planning_year = (project or {}).get("planning_year")
    numeric_years = {int(float(value)) for value in years if value.replace(".", "", 1).isdigit()}
    if planning_year and numeric_years and numeric_years != {int(planning_year)}:
        issues.append({"code": "wrong_planning_year", "message": "File year differs from project planning year.",
                       "file_years": sorted(numeric_years), "project_year": int(planning_year)})
    accepted_scopes = {"project", str((project or {}).get("basin_or_irrigation_area") or "").strip().casefold(),
                       str((project or {}).get("province_or_region") or "").strip().casefold()}
    wrong_scopes = [value for value in scopes if value.casefold() not in accepted_scopes]
    if wrong_scopes:
        issues.append({"code": "wrong_geographic_scope", "message": "File scope differs from project scope.",
                       "file_scopes": wrong_scopes})

    hard_review = any(issue["code"] in {"wrong_planning_year", "wrong_geographic_scope",
                                        "explicit_year_not_parseable",
                                        "required_numeric_field_not_parseable"} for issue in issues)
    if len(tied) > 1 or best["ambiguous"] or sheet_ambiguity:
        state = "AMBIGUOUS"
    elif missing or hard_review:
        state = "REVIEW_REQUIRED"
    else:
        state = "AUTO_MATCHED"
    return {
        **base, "state": state, "detected_data_type": best["data_type"],
        "detected_domain": LABELS.get(best["data_type"], best["data_type"]),
        "mapping": best["mapping"], "missing_required_fields": missing,
        "ambiguous_mapping": best["ambiguous"], "issues": issues,
        "rows": len(rows), "columns": columns, "sheet_names": sheets,
        "selected_sheet": options.get("sheet"), "value_types": _value_types(rows, columns),
        "preview_rows": [row.get("values", {}) for row in rows[:5]],
        "year": years, "scope": scopes, "authority": authorities,
        "confidence": round(min(1.0, best["coverage"] + (0.1 if best["hint"] else 0.0)), 3),
        "evidence": {"filename_hint": best["hint"], "required_matched": best["matched"],
                     "required_total": len(best["required"]), "header_count": len(columns),
                     "sheet_selection": sheet_evidence},
    }
