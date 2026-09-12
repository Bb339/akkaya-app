"""Bounded in-memory parsing; raw uploads are never saved as executable files."""
import csv
import io
import json
import zipfile
from xml.etree.ElementTree import ParseError as XmlParseError
from datetime import date, datetime
from pathlib import PurePath
from typing import Any
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from kds.config import MAX_ROWS, MAX_COLUMNS, MAX_EXPANDED_XLSX_BYTES


class ParseError(ValueError):
    pass


def table(headers: list[Any], records: list[list[Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    columns = [str(v).strip() if v is not None else "" for v in headers]
    if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
        raise ParseError("Headers must be non-empty and unique.")
    if len(columns) > MAX_COLUMNS or len(records) > MAX_ROWS:
        raise ParseError("Table exceeds row/column limit.")
    rows = []
    for index, record in enumerate(records, 2):
        if len(record) != len(columns):
            raise ParseError(f"Row {index}: cell count does not match headers.")
        if any(v not in (None, "") for v in record):
            rows.append({"line": index, "values": dict(zip(columns, record))})
    if not rows:
        raise ParseError("No data rows found.")
    return columns, rows


def parse_csv(content: bytes, options: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        text = content.decode("utf-8-sig")
        delimiter = options.get("delimiter")
        if delimiter is not None and delimiter not in (",", ";", "\t"):
            raise ParseError("Delimiter must be comma, semicolon or tab.")
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        records = []
        for row in reader:
            records.append(row)
            if len(records) > MAX_ROWS + 1:
                raise ParseError("Row limit exceeded.")
        if not records:
            raise ParseError("Empty CSV.")
        return table(records[0], records[1:])
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ParseError(f"Invalid UTF-8 CSV: {exc}") from exc


def parse_xlsx(content: bytes, options: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_XLSX_BYTES:
                raise ParseError("Expanded XLSX exceeds limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet_name = options.get("sheet") or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise ParseError("Selected sheet not found.")
            sheet = workbook[sheet_name]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ParseError("Worksheet exceeds row/column limit.")
            records = []
            for cells in sheet.iter_rows():
                if len(records) >= MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                    raise ParseError("Worksheet exceeds row/column limit.")
                if any(c.data_type == "f" for c in cells):
                    raise ParseError("Formula cells are not accepted; upload explicit values.")
                records.append([c.value.isoformat()[:10] if isinstance(c.value, (date, datetime)) else c.value for c in cells])
            if not records:
                raise ParseError("Empty worksheet.")
            return table(records[0], records[1:])
        finally:
            workbook.close()
    except (zipfile.BadZipFile, InvalidFileException, KeyError, TypeError, XmlParseError) as exc:
        raise ParseError(f"Invalid XLSX: {exc}") from exc


def parse_geojson(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    def invalid_constant(value: str) -> None:
        raise ParseError(f"Non-finite JSON constant is not allowed: {value}")
    try:
        obj = json.loads(content.decode("utf-8-sig"), parse_constant=invalid_constant)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ParseError("Malformed GeoJSON.") from exc
    if not isinstance(obj, dict) or obj.get("type") != "FeatureCollection" or not isinstance(obj.get("features"), list):
        raise ParseError("Expected a GeoJSON FeatureCollection.")
    if not 0 < len(obj["features"]) <= MAX_ROWS:
        raise ParseError("Empty collection or feature limit exceeded.")
    if obj.get("crs"):
        raise ParseError("Use WGS84 longitude/latitude without a legacy CRS declaration.")
    rows, columns = [], []
    for index, feature in enumerate(obj["features"], 1):
        if not isinstance(feature, dict) or feature.get("type") != "Feature" or not isinstance(feature.get("properties"), dict):
            raise ParseError(f"Feature {index}: invalid Feature/properties.")
        values = dict(feature["properties"])
        values["geometry"] = feature.get("geometry")
        for key in values:
            if key not in columns:
                columns.append(key)
        rows.append({"line": index, "values": values})
    if len(columns) > MAX_COLUMNS:
        raise ParseError("Too many properties.")
    return columns, rows


def parse(content: bytes, filename: str, options: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]]]:
    extension = PurePath(filename).suffix.lower()
    if extension == ".csv":
        return parse_csv(content, options)
    if extension == ".xlsx":
        return parse_xlsx(content, options)
    if extension == ".geojson":
        return parse_geojson(content)
    raise ParseError("Allowed extensions: .csv, .xlsx, .geojson.")
