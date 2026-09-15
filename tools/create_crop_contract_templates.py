"""Create deterministic synthetic templates for Phase 7 import contracts."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "data_templates"


def create(name, headers, row):
    book = Workbook()
    sheet = book.active
    sheet.title = "SYNTHETIC_not_official"
    sheet.append(headers)
    sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(36, max(12, max(len(str(c.value or "")) for c in column) + 2))
    book.properties.title = "SYNTHETIC / not_official crop data contract template"
    book.properties.subject = "Example values are synthetic and are not Akkaya observations."
    book.save(OUTPUT / name)


def main():
    create(
        "crop_water_parameters_template.xlsx",
        ["crop_identity", "applicable_year", "geographic_scope", "kc_ini", "kc_mid", "kc_end",
         "stage_value_mode", "p_ini", "p_dev", "p_mid", "p_late", "authority_class", "source",
         "source_reference", "notes"],
        ["TURP (KIRMIZI)", 2024, "synthetic_project", 0.2, 0.9, 0.5, "FRACTIONS", 0.2, 0.3,
         0.3, 0.2, "ASSUMED", "synthetic_test_fixture", "not_official", "SYNTHETIC example; not Akkaya data"],
    )
    create(
        "crop_phenology_template.xlsx",
        ["crop_identity", "applicable_year", "geographic_scope", "season", "mode", "planting_date",
         "harvest_date", "planting_window_start", "planting_window_end", "harvest_window_start",
         "harvest_window_end", "authority_class", "source", "source_reference", "notes"],
        ["TURP (KIRMIZI)", 2024, "synthetic_project", "PRIMARY", "YEAR_SPECIFIC", "2024-04-15",
         "2024-09-15", "", "", "", "", "ASSUMED", "synthetic_test_fixture", "not_official",
         "SYNTHETIC example; not Akkaya data"],
    )


if __name__ == "__main__":
    main()
