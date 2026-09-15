"""Phase 7 crop identity, parameter and phenology contract regressions."""
from __future__ import annotations

import hashlib
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import app
from kds.adapters.akkaya_demo import build_demo
from kds.application.import_service import ImportService
from kds.application.project_service import ProjectService
from kds.application.readiness import readiness
from kds.data.project_store import FileProjectStore
from kds.data.repositories import ConflictError
from kds.domain.crop_parameters import (
    AMBIGUOUS_GRANULARITY_RELATIONS, LEGACY_DIRECT_FALLBACK_CROPS,
    PARAMETER_ONLY_IDENTITIES, REVIEWED_IDENTITY_RELATIONS,
    textual_equivalence_key,
)
from tools.scientific_audit.crop_identity_phenology_phase7 import generate as generate_phase7_evidence


@pytest.fixture
def services(tmp_path):
    store = FileProjectStore(tmp_path / "projects")
    projects = ProjectService(store)
    for project_id in ("one", "two"):
        projects.create(dict(id=project_id, name="Synthetic crop contract", planning_year=2024,
                             annual_water_budget=1000, water_budget_unit="m3",
                             data_source_notes="synthetic_test_fixture"))
        def add_crops(document):
            document["crops"] = [
                {"name": name, "kc_initial": None, "kc_mid": None, "kc_end": None,
                 "stage_initial_days": None, "stage_development_days": None,
                 "stage_mid_days": None, "stage_late_days": None,
                 "confidence_level": "synthetic", "crop_group": ""}
                for name in ("TURP (KIRMIZI)", "KAYSI", "NEKTAR", "KABAK (BAL)",
                             "KABAK (ÇEREZLİK)", "KABAK (SAKIZ)", "SOĞAN (TAZE)")
            ]
        store.update(project_id, add_crops)
    return store, ImportService(store)


def upload(service, kind, content, project="one"):
    return service.upload(project, kind, f"{kind}.csv", content.encode("utf-8"), {})


def confirm(service, batch, project="one", **kwargs):
    return service.confirm(project, batch["id"], acknowledge_warnings=True, **kwargs)


def water_csv(*, crop="TURP (KIRMIZI)", kc_mid=0.9, stages="0.2,0.3,0.3,0.2",
              mode="FRACTIONS", source="synthetic_test_fixture", reference="not_official"):
    return ("crop_identity,applicable_year,geographic_scope,kc_ini,kc_mid,kc_end,stage_value_mode,"
            "p_ini,p_dev,p_mid,p_late,authority_class,source,source_reference,notes\n"
            f"{crop},2024,synthetic_project,0.2,{kc_mid},0.5,{mode},{stages},ASSUMED,{source},{reference},SYNTHETIC\n")


def phenology_csv(*, crop="TURP (KIRMIZI)", planting="2024-04-15", harvest="2024-09-15",
                  year=2024, source="synthetic_test_fixture", reference="not_official"):
    return ("crop_identity,applicable_year,geographic_scope,season,mode,planting_date,harvest_date,"
            "planting_window_start,planting_window_end,harvest_window_start,harvest_window_end,"
            "authority_class,source,source_reference,notes\n"
            f"{crop},{year},synthetic_project,PRIMARY,YEAR_SPECIFIC,{planting},{harvest},,,,,"
            f"ASSUMED,{source},{reference},SYNTHETIC\n")


def test_reviewed_relations_are_explicit_and_unicode_scope_is_narrow():
    assert len(REVIEWED_IDENTITY_RELATIONS) == 9
    assert len(AMBIGUOUS_GRANULARITY_RELATIONS) == 3
    assert len(LEGACY_DIRECT_FALLBACK_CROPS) == 13
    assert len(PARAMETER_ONLY_IDENTITIES) == 14
    relations = {(r.runtime_crop_id, r.parameter_crop_id) for r in REVIEWED_IDENTITY_RELATIONS}
    assert ("KIMYON", "KI\u0307MYON") in relations
    assert ("KAYSI", "KAYISI") in relations and ("NEKTAR", "NEKTARIN") in relations
    assert textual_equivalence_key("KİMYON") == textual_equivalence_key("Ki\u0307myon")
    assert textual_equivalence_key("SALÇALIK DOMATES") != textual_equivalence_key("DOMATES SALÇALIK")


def test_water_parameter_import_validation_and_ambiguity(services):
    _, service = services
    ready = upload(service, "crop_water_parameters", water_csv())
    assert ready["status"] == "ready" and ready["display_label"] == "SYNTHETIC / not_official"
    record = ready["normalized_preview"][0]
    assert record["resolution_status"] == "EXACT" and record["verified_for_pilot"] is False
    assert ready["replacement_preview"]["engine_connected"] is False
    assert upload(service, "crop_water_parameters", water_csv(crop="UNKNOWN"))["status"] == "invalid"
    assert upload(service, "crop_water_parameters", water_csv(crop="KABAK"))["status"] == "invalid"
    assert upload(service, "crop_water_parameters", water_csv(kc_mid=-1))["status"] == "invalid"
    assert upload(service, "crop_water_parameters", water_csv(stages="0.2,0.2,0.2,0.2"))["status"] == "invalid"
    assert upload(service, "crop_water_parameters", water_csv(source=""))["status"] == "invalid"


def test_days_and_reviewed_parameter_identity_are_preserved(services):
    _, service = services
    days = upload(service, "crop_water_parameters", water_csv(mode="DAYS", stages="20,30,40,20"))
    assert days["status"] == "ready"
    assert days["normalized_preview"][0]["stage_total"] == 110
    alias = upload(service, "crop_water_parameters", water_csv(crop="Kayısı"))
    row = alias["normalized_preview"][0]
    assert row["runtime_crop_id"] == "KAYSI" and row["parameter_crop_id"] == "KAYISI"
    assert row["resolution_status"] == "REVIEWED_ALIAS"


@pytest.mark.parametrize("content", [
    phenology_csv(planting="not-a-date"),
    phenology_csv(planting="2024-10-01", harvest="2024-09-01"),
    phenology_csv(planting="2023-04-01"),
    phenology_csv(year=2023, planting="2023-04-01", harvest="2023-09-01"),
    phenology_csv(crop="UNKNOWN"),
    phenology_csv(reference=""),
])
def test_invalid_phenology_is_rejected(services, content):
    _, service = services
    assert upload(service, "crop_phenology", content)["status"] == "invalid"


def test_phenology_modes_and_no_activation_before_confirm(services):
    store, service = services
    year_specific = upload(service, "crop_phenology", phenology_csv())
    assert year_specific["status"] == "ready"
    assert store.get("one")["crop_parameter_data"]["active"] == {}
    confirm(service, year_specific)
    active = store.get("one")["crop_parameter_data"]["active"]
    assert active["crop_phenology|2024|synthetic_project"].startswith("crop-parameter-")
    window = ("crop_identity,applicable_year,geographic_scope,season,mode,planting_window_start,"
              "planting_window_end,harvest_window_start,harvest_window_end,authority_class,source,source_reference\n"
              "TURP (KIRMIZI),2024,synthetic_project,PRIMARY,CLIMATOLOGICAL_WINDOW,04-10,04-25,09-01,09-20,"
              "ASSUMED,synthetic_test_fixture,not_official\n")
    assert upload(service, "crop_phenology", window)["status"] == "ready"


def test_versioning_confirmation_reanalysis_and_project_isolation(services):
    store, service = services
    v1 = upload(service, "crop_water_parameters", water_csv(kc_mid=0.9)); confirm(service, v1)
    old_id = store.get("one")["crop_parameter_data"]["active"]["crop_water_parameters|2024|synthetic_project"]
    v2 = upload(service, "crop_water_parameters", water_csv(kc_mid=1.0))
    assert v2["replacement_preview"]["requires_reanalysis"] is True
    assert store.get("one")["crop_parameter_data"]["active"]["crop_water_parameters|2024|synthetic_project"] == old_id
    confirm(service, v2)
    document = store.get("one")
    new_id = document["crop_parameter_data"]["active"]["crop_water_parameters|2024|synthetic_project"]
    assert new_id != old_id
    assert document["crop_parameter_data"]["datasets"][old_id]["status"] == "inactive"
    assert document["crop_parameter_data"]["datasets"][old_id]["superseded_by"] == new_id
    assert document["crop_parameter_data"]["datasets"][new_id]["version"] == 2
    assert document["crop_parameter_data"]["reanalysis"]["requires_reanalysis"] is True
    assert document["crop_parameter_data"]["reanalysis"]["automatic_recompute"] is False
    assert store.get("two")["crop_parameter_data"]["datasets"] == {}


def test_phenology_replacement_retains_version_history(services):
    store, service = services
    v1 = upload(service, "crop_phenology", phenology_csv(harvest="2024-09-15"))
    confirm(service, v1)
    pointer = "crop_phenology|2024|synthetic_project"
    old_id = store.get("one")["crop_parameter_data"]["active"][pointer]
    v2 = upload(service, "crop_phenology", phenology_csv(harvest="2024-09-20"))
    assert store.get("one")["crop_parameter_data"]["active"][pointer] == old_id
    assert v2["replacement_preview"]["changes"][0]["old_values"]["harvest_date"] == "2024-09-15"
    assert v2["replacement_preview"]["changes"][0]["new_values"]["harvest_date"] == "2024-09-20"
    confirm(service, v2)
    document = store.get("one")
    new_id = document["crop_parameter_data"]["active"][pointer]
    assert document["crop_parameter_data"]["datasets"][old_id]["status"] == "inactive"
    assert document["crop_parameter_data"]["datasets"][old_id]["superseded_by"] == new_id
    assert document["crop_parameter_data"]["datasets"][new_id]["supersedes"] == old_id
    assert document["crop_parameter_data"]["datasets"][new_id]["version"] == 2


def test_current_akkaya_readiness_is_honest_and_engine_disconnected():
    document = build_demo(app.DATA_DIR)
    report = readiness(document)
    params = report["crop_parameter_readiness"]
    assert (params["runtime_crop_count"], params["exact_parameter_count"], params["reviewed_alias_count"]) == (58, 45, 9)
    assert (params["legacy_fallback_count"], params["ambiguous_count"], params["missing_count"]) == (13, 3, 1)
    assert params["identity_resolved_count"] == 54 and params["verified_parameter_count"] == 0
    assert report["phenology_readiness"]["verified_complete_count"] == 0
    assert report["phenology_readiness"]["missing_count"] == 58
    assert report["pilot_readiness"]["status"] == "PILOT_DATA_NOT_READY"
    assert report["pilot_readiness"]["engine_connected"] is False
    assert any(i["code"] == "economic_completeness" and "DUT" in i["message"] for i in report["scenarios"]["S1"]["issues"])


@pytest.mark.parametrize("kind,filename", [
    ("crop_water_parameters", "crop_water_parameters_template.xlsx"),
    ("crop_phenology", "crop_phenology_template.xlsx"),
])
def test_templates_are_formula_free_visible_synthetic_and_importable(services, kind, filename):
    _, service = services
    path = Path(__file__).resolve().parents[1] / "docs" / "data_templates" / filename
    book = load_workbook(path, data_only=False)
    assert sum(cell.data_type == "f" for sheet in book for row in sheet.iter_rows() for cell in row) == 0
    text = " ".join(str(cell.value or "") for sheet in book for row in sheet.iter_rows() for cell in row).lower()
    assert "synthetic" in text and "not_official" in text
    book.close()
    batch = service.upload("one", kind, filename, path.read_bytes(), {})
    assert batch["status"] == "ready" and batch["display_label"] == "SYNTHETIC / not_official"


def test_phase7_scientific_source_guard():
    root = Path(__file__).resolve().parents[1]
    changed = subprocess.run([
        "git", "diff", "--name-only", "38f8c89d4ab14de6868dcd84f017e6e5378353aa", "--",
        "app.py", "kds/science", "kds/application/optimization.py", "data",
    ], cwd=root, check=True, text=True, capture_output=True).stdout.strip()
    assert changed == ""
    assert hashlib.sha256((root / "data/excel_derived/combined_parcel_candidate_matrix_2024.csv").read_bytes()).hexdigest()


def test_phase7_evidence_is_deterministic_and_matches_contract(tmp_path):
    output = tmp_path / "evidence"
    result = generate_phase7_evidence(output)
    first = {path.name: path.read_bytes() for path in output.iterdir()}
    assert result["resolution_counts"] == {
        "EXACT": 45, "REVIEWED_ALIAS": 9, "AMBIGUOUS": 3, "MISSING": 1,
    }
    assert result["legacy_fallback_count"] == 13
    assert result["phenology_verified_count"] == 0
    generate_phase7_evidence(output)
    assert first == {path.name: path.read_bytes() for path in output.iterdir()}
